"""Load the CMS Transparency in Coverage PUF into the database.

Shape of the source workbook:

* Four sheets - a disclaimer, then Individual QHP, Individual SADP and SHOP.
  All three data sheets share identical headers.
* Headers sit on the THIRD row. Rows one and two are a title banner and a blank.
* Issuer-level figures are repeated on every plan row belonging to that issuer,
  which is why `sql/metrics.sql` collapses them with MAX rather than SUM.
* Every numeric column is typed as text. See `app.coerce`.

Columns are matched by name, not position, and the names are normalised before
matching. CMS ships genuine typos in its own headers - "Enrolle" for enrollee,
"Proceduce" for procedure - so both the misspelling and the correct spelling are
accepted. A column order change or a spelling fix upstream then costs nothing;
positional indexing would silently load the wrong data into the wrong field.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from app.coerce import to_int, to_text

HEADER_ROW = 3  # 1-indexed, as openpyxl counts

# Sheet names carry the plan year and are awkward to read in output. The market
# is the important part, and it is the axis every comparison is segmented on:
# dental plans deny at structurally different rates from medical ones, so
# ranking them together just rediscovers "dental is not medical".
MARKETS = {
    "Transparency 2025 - Ind QHP": "Individual QHP",
    "Transparency 2025 - Ind SADP": "Individual SADP",
    "Transparency 2025 - SHOP": "SHOP",
}


def market_name(sheet_name: str) -> str:
    return MARKETS.get(sheet_name, sheet_name)
DISCLAIMER_SHEET = "PUF Data Disclaimer"

TEXT_COLUMNS = {
    "State": "state",
    "Issuer_Name": "issuer_name",
    "Issuer_ID": "issuer_id",
    "Plan_ID": "plan_id",
    "Plan_Type": "plan_type",
    "Metal_Level": "metal_level",
}

INT_COLUMNS = {
    "Issuer_Claims_Received_In_Network": "issuer_claims_received_in_network",
    "Issuer_Claims_Denied_In_Network": "issuer_claims_denied_in_network",
    "Issuer_Claims_Received_Out_of_Network": "issuer_claims_received_out_network",
    "Issuer_Claims_Denied_Out_of_Network": "issuer_claims_denied_out_network",
    "Plan_Number_Claims_Received_In_Network": "plan_claims_received_in_network",
    "Plan_Number_Claims_Denied_In_Network": "plan_claims_denied_in_network",
    "Plan_Number_Claims_Denied_Referral_Required": "denied_referral_required",
    "Plan_Number_Claims_Denied_Due_To_Out_Of_Network": "denied_out_of_network",
    "Plan_Number_Claims_Denied_Services_Excluded": "denied_service_excluded",
    "Plan_Number_Claims_Denied_Not_Medically_Necessary_Excluding_Behavioral_Health":
        "denied_not_medically_necessary",
    "Plan_Number_Claims_Denied_Not_Medically_Necessary_Behavioral_Health_Only":
        "denied_behavioral_health",
    # "Enrolle" is CMS's typo, not ours.
    "Plan_Number_Claims_Denied_Due_To_Enrolle_Benefit_Limit_Reached":
        "denied_benefit_limit_reached",
    "Plan_Number_Claims_Denied_Due_To_Enrollee_Benefit_Limit_Reached":
        "denied_benefit_limit_reached",
    "Plan_Number_Claims_Denied_Due_To_Member_Not_Covered": "denied_member_not_covered",
    # "Proceduce" likewise.
    "Plan_Number_Claims_Denied_Due_To_Investigational_Experimental_Cosmetic_Proceduce":
        "denied_experimental",
    "Plan_Number_Claims_Denied_Due_To_Investigational_Experimental_Cosmetic_Procedure":
        "denied_experimental",
    "Plan_Number_Claims_Denied_Due_To_Administrative_Reason": "denied_administrative",
    "Plan_Number_Claims_Denied_Other": "denied_other",
    "Issuer_Internal_Appeals_Filed": "internal_appeals_filed",
    "Issuer_Number_Internal_Appeals_Overturned": "internal_appeals_overturned",
    "Issuer_External_Appeals_Filed": "external_appeals_filed",
    "Issuer_Number_External_Appeals_Overturned": "external_appeals_overturned",
    "Average Monthly Enrollment": "average_monthly_enrollment",
}

REQUIRED = ("state", "issuer_id", "issuer_name")


def normalise(header: object) -> str:
    """Lowercase and strip everything that is not a letter or digit."""
    return re.sub(r"[^a-z0-9]", "", str(header or "").lower())


TEXT_LOOKUP = {normalise(k): v for k, v in TEXT_COLUMNS.items()}
INT_LOOKUP = {normalise(k): v for k, v in INT_COLUMNS.items()}


@dataclass
class LoadReport:
    rows_loaded: int = 0
    blank_rows: int = 0
    malformed_rows: int = 0
    per_market: dict[str, int] = field(default_factory=dict)
    unmapped_headers: list[str] = field(default_factory=list)

    @property
    def rows_skipped(self) -> int:
        return self.blank_rows + self.malformed_rows


def is_blank(values: list) -> bool:
    """A row with nothing in it at all. Spreadsheets are full of these."""
    return not any(str(v).strip() for v in values if v is not None)


def map_row(headers: list, values: list) -> dict | None:
    """Turn one spreadsheet row into a database row, or None if unusable.

    A row without a state and an issuer cannot be attributed to anyone, so it is
    dropped rather than stored as an anonymous set of numbers.
    """
    record: dict = {}
    for header, value in zip(headers, values, strict=False):
        key = normalise(header)
        if key in TEXT_LOOKUP:
            record[TEXT_LOOKUP[key]] = to_text(value)
        elif key in INT_LOOKUP:
            record[INT_LOOKUP[key]] = to_int(value)

    if any(not record.get(name) for name in REQUIRED):
        return None
    return record


def unmapped(headers: list) -> list[str]:
    """Headers we ignore. Surfaced so a schema change upstream is visible."""
    return [
        str(h) for h in headers
        if h and normalise(h) not in TEXT_LOOKUP and normalise(h) not in INT_LOOKUP
    ]


def read_sheets(path: str | Path):
    """Yield (sheet_name, headers, row_values) for each data sheet."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet_name in workbook.sheetnames:
            if sheet_name == DISCLAIMER_SHEET:
                continue
            sheet = workbook[sheet_name]
            rows = sheet.iter_rows(values_only=True)
            headers = None
            for index, row in enumerate(rows, start=1):
                if index < HEADER_ROW:
                    continue
                if index == HEADER_ROW:
                    headers = list(row)
                    continue
                yield sheet_name, headers, list(row)
    finally:
        workbook.close()


def load(path: str | Path, connection) -> LoadReport:
    report = LoadReport()
    seen_headers = False
    batch: list[tuple[str, dict]] = []

    for sheet_name, headers, values in read_sheets(path):
        if not seen_headers:
            report.unmapped_headers = unmapped(headers)
            seen_headers = True

        if is_blank(values):
            # Counted separately from malformed rows on purpose. If these were
            # lumped together, one genuinely broken row would hide among the
            # thousand empty ones and nobody would ever notice.
            report.blank_rows += 1
            continue

        record = map_row(headers, values)
        if record is None:
            report.malformed_rows += 1
            continue

        market = market_name(sheet_name)
        batch.append((market, record))
        report.per_market[market] = report.per_market.get(market, 0) + 1

    if batch:
        columns = ["plan_market"] + sorted(batch[0][1])
        placeholders = ", ".join("?" for _ in columns)
        connection.executemany(
            f"INSERT INTO plan_claims ({', '.join(columns)}) VALUES ({placeholders})",
            [(sheet, *[rec.get(c) for c in columns[1:]]) for sheet, rec in batch],
        )
        connection.commit()
        report.rows_loaded = len(batch)

    return report
